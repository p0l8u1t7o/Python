import os
import xml.etree.ElementTree as ET
from opencc import OpenCC
from openpyxl import Workbook
from typing import Tuple, List


#============================================================================================================================================#
# 讀取 TXT 檔案內容
def read_txt_file(filepath: str) -> str:
    if not filepath.strip():
        return ""
    
    if not os.path.exists(filepath):
        return ""

    with open(filepath, 'r', encoding='utf-8') as file:
        content = file.read()
    
    return content


#============================================================================================================================================#
# 解析 SOAP XML 檔案，並建立 Field 的 Tuple（key, value）列表

def parse_soap_field_map(content: str) -> Tuple[List[str], List[str]]:
    try:
        root = ET.fromstring(content)
        fields = root.findall(".//Field")

        name_value_pairs = []
        for field in fields:
            name = field.get('name')
            value = field.get('value') or ''
            if name is not None:
                name_value_pairs.append((name, value))

        # 根據名稱反序排序
        name_value_pairs.sort(key=lambda x: x[0], reverse=True)

        # 分離出排序後的 names 和 values
        names = [pair[0] for pair in name_value_pairs]
        values = [pair[1] for pair in name_value_pairs]

        return names, values

    except ET.ParseError as e:
        print("XML parsing error:", e)
        return [], []

#============================================================================================================================================#
# 將 Field 的 Tuple（key, value） 轉換為中文

def translate_to_traditional_chinese_tuple(data: Tuple[List[str], List[str]]) -> Tuple[List[str], List[str]]:
    # 建立 OpenCC 轉換器，指定轉換模式：s2t（簡體轉繁體）
    cc = OpenCC('s2t')

    # 解構傳入的 tuple，分別取出 keys 和 values 兩個 list
    keys, values = data

    # 將 values 中的每個元素轉換成繁體中文（如果是字串的話）
    converted_values = [cc.convert(v) if isinstance(v, str) else v for v in values]

    # 回傳原本的 keys 和轉換後的 values 組成的 tuple
    return keys, converted_values


#============================================================================================================================================#
# 將 Tuple（key, value） 寫入 Excel 檔案

def tuple_to_excel(data: Tuple[List[str], List[str]], filename: str):
    if not filename.strip():
        return

    keys, values = data

    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"  # 設定工作表名稱為 sheet1

    # 寫入 Key（第一列）
    for col_index, key in enumerate(keys, start=1):
        ws.cell(row=1, column=col_index, value=key)

    # 寫入 Value（第二列）
    for col_index, value in enumerate(values, start=1):
        ws.cell(row=2, column=col_index, value=value)

    wb.save(filename)
    wb.close()

#============================================================================================================================================#
# 使用 OpenCC 進行簡體中文到繁體中文

def translate_to_traditional_chinese_string(simplified_text: str) -> str:

    # 建立一個 OpenCC 轉換器，使用 s2t 模式（簡體轉繁體）
    cc = OpenCC('s2t')

    # 轉換成繁體中文
    traditional_text = cc.convert(simplified_text)

    return traditional_text