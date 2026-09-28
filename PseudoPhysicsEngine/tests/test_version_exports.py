"""版本匯出一致性：匯出指定版本時只能讀該版本的快照，不得混入工作區或其他版本。"""

from __future__ import annotations

import csv
import hashlib
import json
from datetime import date
from pathlib import Path

import ezdxf
import pytest
from docx import Document

from cellforge import derived
from cellforge.agents.astra import refresh_theme_local
from cellforge.build.pipeline import build_project
from cellforge.exports import export_project
from cellforge.project import create_project
from cellforge.versioning import VersionNotFoundError
from cellforge.yamlio import dump_yaml, load_yaml

V1_NAME = "原始案名"
V2_NAME = "改版案名"
V2_ASSUMPTION = "v2 才加入的假設：治具高度改為可調"
V2_CHECKLIST = "v2 才有的 checklist 對照內容"
V2_MODULE = "v2_extra_stand"
KINDS = ["step", "dxf", "bom", "report", "html"]


def _make_project(root: Path) -> Path:
    return create_project(
        root / "consistency",
        {"name": V1_NAME, "customer": "Getac", "created": date.today().isoformat()},
        seed_example="getac_qc",
    )


def _mutate_workspace(project: Path) -> None:
    data = load_yaml(project / "project.yaml")
    data["name"] = V2_NAME
    dump_yaml(project / "project.yaml", data)
    assumptions = load_yaml(project / "analysis" / "assumptions.yaml")
    assumptions["assumptions"].append(
        {"id": "A-V2-ONLY", "text": V2_ASSUMPTION, "basis": "版本一致性測試"}
    )
    dump_yaml(project / "analysis" / "assumptions.yaml", assumptions)
    (project / "analysis" / "checklist_map.md").write_text(V2_CHECKLIST, encoding="utf-8")
    cell = load_yaml(project / "cell.yaml")
    cell["machines"][0]["modules"].append(
        {
            "id": V2_MODULE,
            "part": "library/fixture_stand.py",
            "params": {},
            "pose": {"xyz": [0, -2500, 0], "rpy_deg": [0, 0, 0], "trust": "inferred"},
        }
    )
    dump_yaml(project / "cell.yaml", cell)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _outputs(report: dict) -> dict[str, Path]:
    directory = Path(report["directory"])
    return {Path(name).suffix: directory / name for name in report["files"]}


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def _bom_modules(path: Path) -> set[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return {row["module_id"] for row in csv.DictReader(handle)}


def _dxf_labels(path: Path) -> set[str]:
    return {entity.dxf.text for entity in ezdxf.readfile(path).modelspace().query("TEXT")}


def test_export_v1_after_v2_reads_only_v1_snapshot(tmp_path: Path):
    project = _make_project(tmp_path)
    assert build_project(project, "L0")["version"] == 1
    v1_modules = {
        module["id"]
        for machine in load_yaml(project / "cell.yaml")["machines"]
        for module in machine["modules"]
    }
    _mutate_workspace(project)
    assert build_project(project, "L0")["version"] == 2

    v1 = export_project(project, KINDS, "v1")
    assert v1["version"] == "v1"
    files = _outputs(v1)
    assert all(V2_NAME not in path.name for path in files.values())
    assert files[".step"].read_bytes() == (project / ".cellforge/v1/scene.step").read_bytes()
    assert _bom_modules(files[".csv"]) == v1_modules
    assert V2_MODULE not in _dxf_labels(files[".dxf"])
    report_text = _docx_text(files[".docx"])
    assert V1_NAME in report_text
    for marker in (V2_NAME, V2_ASSUMPTION, V2_CHECKLIST):
        assert marker not in report_text
    html = files[".html"].read_text("utf-8")
    assert f"<title>{V1_NAME}</title>" in html
    assert V2_NAME not in html

    # 因果對照：同一批匯出器讀 v2 時必須看得到修改，證明上面的斷言有鑑別力。
    v2 = export_project(project, KINDS, "v2")
    files = _outputs(v2)
    assert files[".step"].read_bytes() == (project / ".cellforge/v2/scene.step").read_bytes()
    assert files[".step"].read_bytes() != (project / ".cellforge/v1/scene.step").read_bytes()
    assert _bom_modules(files[".csv"]) == v1_modules | {V2_MODULE}
    assert V2_MODULE in _dxf_labels(files[".dxf"])
    report_text = _docx_text(files[".docx"])
    for marker in (V2_NAME, V2_ASSUMPTION, V2_CHECKLIST):
        assert marker in report_text
    assert f"<title>{V2_NAME}</title>" in files[".html"].read_text("utf-8")

    latest = export_project(project, ["bom"])
    assert latest["version"] == "v2"


def test_snapshot_missing_required_source_is_reported_not_read_from_workspace(tmp_path: Path):
    project = _make_project(tmp_path)
    build_project(project, "L0")
    # 模擬舊版本缺件；工作區的 cell.yaml 仍在，匯出器不得拿它補。
    (project / ".cellforge" / "v1" / "source" / "cell.yaml").unlink()
    assert (project / "cell.yaml").is_file()
    stale = project / "export" / "v1" / f"{V1_NAME}_v1_bom.csv"
    stale.parent.mkdir(parents=True, exist_ok=True)
    stale.write_text("舊匯出殘留", encoding="utf-8")
    report = export_project(project, ["step", "bom", "dxf"], "v1")
    assert report["status"] == "incomplete"
    assert {item["kind"] for item in report["missing"]} == {"bom", "dxf"}
    for item in report["missing"]:
        assert "版本資料不完整" in item["reason"] and "cell.yaml" in item["reason"]
    assert [Path(name).suffix for name in report["files"]] == [".step"]
    assert not stale.exists()
    manifest = json.loads((project / "export" / "v1" / "manifest.json").read_text("utf-8"))
    assert manifest["status"] == "incomplete"
    assert manifest["missing"] == report["missing"]


def test_legacy_snapshot_exports_covered_sources_with_warning(tmp_path: Path):
    project = _make_project(tmp_path)
    build_project(project, "L0")
    # 移除 manifest 模擬 DEV-003 之前的舊版快照；範圍內不存在的 checklist 視為建置時沒有。
    (project / ".cellforge" / "v1" / "manifest.json").unlink()
    (project / "analysis" / "checklist_map.md").write_text(V2_CHECKLIST, encoding="utf-8")
    report = export_project(project, ["bom", "report"], "v1")
    assert report["status"] == "ok"
    assert any("舊版快照" in warning for warning in report["warnings"])
    report_text = _docx_text(_outputs(report)[".docx"])
    assert "舊版快照" in report_text
    assert V2_CHECKLIST not in report_text
    manifest = json.loads((project / "export" / "v1" / "manifest.json").read_text("utf-8"))
    assert manifest["version_complete"] is False


def test_export_records_hashes_and_lists_missing_video(tmp_path: Path):
    project = _make_project(tmp_path)
    build_project(project, "L0")
    # 工作區有「最新」影片，也有舊匯出器複製過來的同名檔；兩者都不能冒充 v1 的影片。
    presentation = project / "presentation"
    presentation.mkdir(exist_ok=True)
    (presentation / "cell_review_1080p.mp4").write_bytes(b"latest workspace video")
    stale = project / "export" / "v1" / f"{V1_NAME}_v1_1080p.mp4"
    stale.parent.mkdir(parents=True, exist_ok=True)
    stale.write_bytes(b"latest workspace video")
    report = export_project(project, [*KINDS, "video"], "v1")
    assert report["status"] == "incomplete"
    assert [item["kind"] for item in report["missing"]] == ["video"]
    assert "尚未產生影片" in report["missing"][0]["reason"]
    assert not stale.exists()
    manifest = json.loads((project / "export" / "v1" / "manifest.json").read_text("utf-8"))
    assert manifest["version_complete"] is True
    assert manifest["version_manifest_sha256"] == _sha256(project / ".cellforge/v1/manifest.json")
    assert {item["kind"] for item in manifest["outputs"]} == set(KINDS)
    for item in manifest["outputs"]:
        assert _sha256(project / "export" / "v1" / item["file"]) == item["sha256"]
    html = next(item for item in manifest["outputs"] if item["kind"] == "html")
    assert html["provenance"]["offline_viewer_bundle_sha256"]


def test_video_export_uses_only_the_requested_versions_video(tmp_path: Path):
    project = _make_project(tmp_path)
    build_project(project, "L0")
    astra = refresh_theme_local(project, "版本影片測試")
    v1_video = project / astra["video"]
    assert v1_video.is_file() and "TEMP" in v1_video.parts
    registered = derived.latest_available(project, "v1", "video")
    assert registered is not None and registered.path == v1_video.resolve()
    assert registered.entry["settings"]["placeholder_only"] is True

    build_project(project, "L0")
    v2 = export_project(project, ["video"], "v2")
    assert v2["status"] == "incomplete" and v2["files"] == []
    v1 = export_project(project, ["video"], "v1")
    assert v1["status"] == "ok"
    assert _outputs(v1)[".mp4"].read_bytes() == v1_video.read_bytes()

    # 暫存影片被清掉後，v1 顯示可重新產生，而不是仍可下載。
    v1_video.unlink()
    cleared = export_project(project, ["video"], "v1")
    assert cleared["status"] == "incomplete"
    assert "可重新產生" in cleared["missing"][0]["reason"]
    assert not _outputs(v1)[".mp4"].exists()


def test_tampered_version_artifact_is_reported_not_exported(tmp_path: Path):
    project = _make_project(tmp_path)
    build_project(project, "L0")
    timeline = project / ".cellforge" / "v1" / "timeline.json"
    timeline.write_text(timeline.read_text("utf-8") + " ", encoding="utf-8")
    report = export_project(project, ["step", "report", "html"], "v1")
    assert [Path(name).suffix for name in report["files"]] == [".step"]
    assert {item["kind"] for item in report["missing"]} == {"report", "html"}
    assert all("雜湊不符" in item["reason"] for item in report["missing"])


@pytest.mark.parametrize("version", ["v9", "../v1", "v1/../v1", "latest"])
def test_export_rejects_unknown_or_unsafe_version(tmp_path: Path, version: str):
    project = _make_project(tmp_path)
    build_project(project, "L0")
    with pytest.raises(VersionNotFoundError, match="版本"):
        export_project(project, ["bom"], version)
