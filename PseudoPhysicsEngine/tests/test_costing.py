"""DEV-009／ACC-05：成本依公式更新、缺價單獨列出、不重複計價。

第一組測試用獨立的小型人工核算案例：期望值在測試內以十進位逐步手算（與平台程式無共用程式碼），
不以平台自己產生的答案作唯一依據。其餘測試驗證增刪設備、改單價、缺價與資料錯誤的因果。
"""

from __future__ import annotations

import copy
import csv
import json
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

from cellforge.build.pipeline import build_project
from cellforge.costing import compute_costing
from cellforge.exports import export_project
from cellforge.project import create_project
from cellforge.schema.costing import Costing
from cellforge.schema.models import Cell
from cellforge.yamlio import dump_yaml, load_yaml
from server.main import create_app

CAM_PRICE = "12345.67"
LENS_USD = "100.25"
USD_RATE = "31.5"
KIT_PRICE = "50000"
PLATE_PRICE = "1500"
BOLT_PRICE = "2.5"
INSTALL_PRICE = "8000"
# 3 × 0.335 = 1.005：明細小計捨入到 2 位（四捨五入）時為 1.01，驗證捨入階段
WASHER_PRICE = "0.335"
WASHERS = 3
LABOR_RATE = "800"
MECH_HOURS = "10"
SOFT_HOURS = "12.5"
OVERHEAD, CONTINGENCY, TAX, MARGIN = "0.05", "0.10", "0.05", "0.2"


def _costing(**changes) -> dict:
    data = {
        "as_of": "2026-09-28",
        "policy": {
            "currency": "TWD",
            "overhead_rate": OVERHEAD,
            "contingency_rate": CONTINGENCY,
            "tax_rate": TAX,
            "rounding": {"line_decimals": 2, "total_decimals": 0, "mode": "half_up"},
            "pricing": {"method": "margin", "rate": MARGIN},
        },
        "exchange_rates": [
            {
                "from": "USD",
                "to": "TWD",
                "rate": USD_RATE,
                "date": "2026-09-01",
                "source": "測試匯率",
            }
        ],
        "items": [
            {
                "id": "cam",
                "kind": "purchased",
                "subsystem": "視覺",
                "name": "相機",
                "unit": "台",
                "price": {"amount": CAM_PRICE},
            },
            {
                "id": "lens",
                "kind": "purchased",
                "subsystem": "視覺",
                "name": "鏡頭",
                "unit": "支",
                "price": {"amount": LENS_USD, "currency": "USD", "range_pct": "0.2"},
            },
            {
                "id": "kit",
                "kind": "module",
                "subsystem": "機構",
                "name": "有單價的組件",
                "unit": "組",
                "price": {"amount": KIT_PRICE},
                "components": [{"item": "screw", "quantity": 4}],
            },
            {
                "id": "screw",
                "kind": "consumable",
                "subsystem": "機構",
                "name": "螺絲",
                "unit": "支",
                "price": {"amount": "3.333"},
            },
            {
                "id": "frame",
                "kind": "fabricated",
                "subsystem": "機構",
                "name": "待報價組件",
                "unit": "組",
                "price": {"amount": None},
                "components": [{"item": "plate", "quantity": 2}, {"item": "bolt", "quantity": 8}],
            },
            {
                "id": "plate",
                "kind": "fabricated",
                "subsystem": "機構",
                "name": "板件",
                "unit": "片",
                "price": {"amount": PLATE_PRICE},
            },
            {
                "id": "bolt",
                "kind": "consumable",
                "subsystem": "機構",
                "name": "螺栓",
                "unit": "支",
                "price": {"amount": BOLT_PRICE},
            },
            {
                "id": "washer",
                "kind": "consumable",
                "subsystem": "機構",
                "name": "墊片",
                "unit": "片",
                "price": {"amount": WASHER_PRICE},
            },
            {
                "id": "install",
                "kind": "service",
                "subsystem": "工程",
                "name": "安裝",
                "unit": "式",
                "price": {"amount": INSTALL_PRICE},
            },
            {
                "id": "plc",
                "kind": "purchased",
                "subsystem": "控制",
                "name": "PLC",
                "unit": "套",
                "price": {"amount": None},
            },
        ],
        "equipment": [
            {
                "id": "EQ-cam",
                "match": {"part": "parts/cam.py"},
                "items": [{"item": "cam", "quantity": 1}, {"item": "lens", "quantity": 1}],
            }
        ],
        "bom": [
            {"id": "B-kit", "item": "kit", "quantity": 1, "for": ["base"]},
            {"id": "B-frame", "item": "frame", "quantity": 1},
            {"id": "B-install", "item": "install", "quantity": 1},
            {"id": "B-plc", "item": "plc", "quantity": 2},
            {"id": "B-washer", "item": "washer", "quantity": WASHERS},
        ],
        "labor": {
            "rates": [{"id": "eng", "rate": LABOR_RATE, "unit": "時"}],
            "tasks": [
                {
                    "id": "L-mech",
                    "category": "mechanical",
                    "name": "機構設計",
                    "quantity": MECH_HOURS,
                    "rate": "eng",
                },
                {
                    "id": "L-soft",
                    "category": "software",
                    "name": "程式",
                    "quantity": SOFT_HOURS,
                    "rate": "eng",
                },
            ],
        },
    }
    for key, value in changes.items():
        data[key] = value
    return data


def _cell(cameras: int = 2) -> Cell:
    modules = [{"id": "base", "part": "library/box.py"}] + [
        {"id": f"cam_{index + 1}", "part": "parts/cam.py"} for index in range(cameras)
    ]
    return Cell.model_validate({"plant_frame": {}, "machines": [{"id": "m", "modules": modules}]})


def _run(data: dict, cameras: int = 2) -> dict:
    return compute_costing(Costing.model_validate(data), _cell(cameras))


def _line(result: dict, item: str) -> dict | None:
    return next((line for line in result["purchase_lines"] if line["item"] == item), None)


def _d(value) -> Decimal:
    return Decimal(str(value))


def _round(value: Decimal, places: int) -> Decimal:
    return value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)


def test_totals_match_an_independent_hand_calculation():
    result = _run(_costing())
    cameras = 2
    # 人工核算：材料 = Σ 數量 × 單價 × 匯率（每列捨入到 2 位）
    cam = _round(cameras * _d(CAM_PRICE), 2)
    lens = _round(cameras * _d(LENS_USD) * _d(USD_RATE), 2)
    kit = _round(_d(KIT_PRICE), 2)  # 組件有單價：螺絲不另計
    plate = _round(2 * _d(PLATE_PRICE), 2)  # 組件待報價：展開子件
    bolt = _round(8 * _d(BOLT_PRICE), 2)
    washer = _round(WASHERS * _d(WASHER_PRICE), 2)
    equipment = cam + lens + kit + plate + bolt + washer
    labor = _round(_d(MECH_HOURS) * _d(LABOR_RATE), 2) + _round(_d(SOFT_HOURS) * _d(LABOR_RATE), 2)
    development = labor + _d(INSTALL_PRICE)
    base = equipment + development
    overhead = _round(_d(OVERHEAD) * base, 0)
    contingency = _round(_d(CONTINGENCY) * (base + overhead), 0)
    cost = _round(base + overhead + contingency, 0)
    budget_tax = _round(_d(TAX) * cost, 0)
    price = _round(cost / (1 - _d(MARGIN)), 0)
    quote_tax = _round(_d(TAX) * price, 0)

    totals = result["totals"]
    expected = {
        "equipment": equipment,
        "labor": labor,
        "development": development,
        "overhead": overhead,
        "contingency": contingency,
        "cost": cost,
        "budget_tax": budget_tax,
        "budget_with_tax": cost + budget_tax,
        "price": price,
        "quote_tax": quote_tax,
        "quote_with_tax": price + quote_tax,
    }
    for key, value in expected.items():
        assert Decimal(totals[key]["value"]) == value, key
    assert Decimal(_line(result, "lens")["amount_high"]) == _round(lens * _d("1.2"), 2)
    assert _line(result, "kit")["components_not_costed"] == ["screw"]
    assert _line(result, "screw") is None
    assert _line(result, "frame") is None and _line(result, "plate") is not None


def test_missing_prices_are_listed_and_never_counted_as_zero():
    result = _run(_costing())
    assert result["complete"] is False
    missing = {entry["id"]: entry for entry in result["missing"]}
    assert set(missing) == {"plc"}
    assert missing["plc"]["reason"] == "待報價" and missing["plc"]["quantity"] == "2"
    assert _line(result, "plc")["amount"] is None
    coverage = result["coverage"]
    assert coverage["priced_lines"] == coverage["total_lines"] - 1
    # 補上單價後，成本增加的正是 數量 × 單價（證明缺價先前沒有被當成 0 以外的值計入）
    priced = _costing()
    next(item for item in priced["items"] if item["id"] == "plc")["price"]["amount"] = "700"
    after = _run(priced)
    assert after["complete"] is True
    delta = Decimal(after["totals"]["equipment"]["value"]) - Decimal(
        result["totals"]["equipment"]["value"]
    )
    assert delta == 2 * Decimal("700")


def test_adding_and_removing_a_camera_updates_quantity_and_cost_without_leftovers():
    two = _run(_costing(), cameras=2)
    three = _run(_costing(), cameras=3)
    assert _line(three, "cam")["quantity"] == "3" and _line(two, "cam")["quantity"] == "2"
    delta = Decimal(three["totals"]["equipment"]["value"]) - Decimal(
        two["totals"]["equipment"]["value"]
    )
    assert delta == _round(_d(CAM_PRICE), 2) + (
        _round(3 * _d(LENS_USD) * _d(USD_RATE), 2) - _round(2 * _d(LENS_USD) * _d(USD_RATE), 2)
    )
    none = _run(_costing(), cameras=0)
    assert _line(none, "cam") is None and _line(none, "lens") is None
    assert any(issue["code"] == "NO_EQUIPMENT" for issue in none["issues"])
    sources = {source["module"] for source in _line(three, "cam")["sources"]}
    assert sources == {"cam_1", "cam_2", "cam_3"}


def test_changing_a_unit_price_changes_the_total_by_quantity_times_the_difference():
    data = _costing()
    before = _run(data)
    changed = copy.deepcopy(data)
    next(item for item in changed["items"] if item["id"] == "cam")["price"]["amount"] = "12445.67"
    after = _run(changed)
    delta = Decimal(after["totals"]["equipment"]["value"]) - Decimal(
        before["totals"]["equipment"]["value"]
    )
    assert delta == 2 * (Decimal("12445.67") - _d(CAM_PRICE))


def test_the_same_item_for_the_same_equipment_is_counted_once_and_reported():
    data = _costing()
    data["bom"].append({"id": "B-cam-extra", "item": "cam", "quantity": 1, "for": ["cam_1"]})
    result = _run(data)
    assert _line(result, "cam")["quantity"] == "2"
    assert any(issue["code"] == "DUPLICATE" for issue in result["issues"])


def test_a_priced_assembly_and_its_component_for_the_same_equipment_is_an_error():
    data = _costing()
    data["bom"].append({"id": "B-screw", "item": "screw", "quantity": 4, "for": ["base"]})
    result = _run(data)
    assert any(issue["code"] == "ASSEMBLY_DOUBLE_COUNT" for issue in result["issues"])


def test_missing_exchange_rate_unit_mismatch_and_expired_quote_are_reported():
    data = _costing()
    lens = next(item for item in data["items"] if item["id"] == "lens")
    lens["price"]["currency"] = "EUR"
    data["bom"].append({"id": "B-bolt", "item": "bolt", "quantity": 3, "unit": "盒"})
    install = next(item for item in data["items"] if item["id"] == "install")
    install["price"]["valid_until"] = "2026-01-31"
    result = _run(data)
    codes = {issue["code"] for issue in result["issues"]}
    assert {"NO_EXCHANGE_RATE", "UNIT_MISMATCH", "QUOTE_EXPIRED"} <= codes
    reasons = {entry["id"]: entry["reason"] for entry in result["missing"]}
    assert reasons["lens"].startswith("缺匯率") and reasons["bolt"].startswith("單位不符")


def test_markup_and_margin_are_different_explicit_formulas():
    markup = _costing()
    markup["policy"]["pricing"] = {"method": "markup", "rate": "0.25"}
    margin = _run(_costing())
    marked = _run(markup)
    cost = Decimal(marked["totals"]["cost"]["value"])
    assert Decimal(marked["totals"]["price"]["value"]) == _round(cost * Decimal("1.25"), 0)
    assert Decimal(margin["totals"]["price"]["value"]) == _round(cost / Decimal("0.8"), 0)
    budget_only = _costing()
    budget_only["policy"]["pricing"] = None
    assert "price" not in _run(budget_only)["totals"]
    with pytest.raises(ValidationError):
        Costing.model_validate({"policy": {"pricing": {"method": "margin", "rate": 1}}})


def test_decimal_inputs_do_not_pick_up_binary_float_error():
    parsed = Costing.model_validate(load_yaml_text("policy: {contingency_rate: 0.1}"))
    assert parsed.policy.contingency_rate == Decimal("0.1")


def load_yaml_text(text: str):
    import io

    from ruamel.yaml import YAML

    return YAML().load(io.StringIO(text))


def test_ssd_example_totals_equal_an_independent_sum_of_its_source_data():
    root = Path(__file__).resolve().parents[1] / "examples" / "ssd_press" / "handwritten"
    raw = load_yaml(root / "costing.yaml")
    cell = load_yaml(root / "cell.yaml")
    modules = [module for machine in cell["machines"] for module in machine["modules"]]
    prices = {item["id"]: _d(item["price"]["amount"]) for item in raw["items"]}
    kinds = {item["id"]: item["kind"] for item in raw["items"]}

    def matched(match):
        for module in modules:
            tool = (module.get("params") or {}).get("tool") or {}
            if all(
                (
                    match.get("part") in (None, module.get("part")),
                    match.get("module") in (None, module["id"]),
                    match.get("tool") in (None, tool.get("part")),
                )
            ):
                yield module

    equipment = services = Decimal("0")
    lines = [
        (entry["item"], _d(entry["quantity"]))
        for rule in raw["equipment"]
        for _module in matched(rule["match"])
        for entry in rule["items"]
    ] + [(line["item"], _d(line["quantity"])) for line in raw["bom"]]
    for item, quantity in lines:
        if kinds[item] == "service":
            services += quantity * prices[item]
        else:
            equipment += quantity * prices[item]
    rates = {rate["id"]: _d(rate["rate"]) for rate in raw["labor"]["rates"]}
    labor = sum(
        (_d(task["quantity"]) * rates[task["rate"]] for task in raw["labor"]["tasks"]),
        Decimal("0"),
    )
    contingency = _d(raw["policy"]["contingency_rate"]) * (equipment + services + labor)
    cost = equipment + services + labor + contingency
    result = compute_costing(Costing.model_validate(raw), Cell.model_validate(cell))
    assert result["complete"] is True and result["issues"] == []
    assert Decimal(result["totals"]["equipment"]["value"]) == equipment
    assert Decimal(result["totals"]["development"]["value"]) == services + labor
    assert Decimal(result["totals"]["cost"]["value"]) == _round(cost, 2)
    tax = _round(_d(raw["policy"]["tax_rate"]) * _round(cost, 2), 2)
    assert Decimal(result["totals"]["budget_with_tax"]["value"]) == _round(cost, 2) + tax


GETAC_COSTING = {
    "as_of": "2026-09-28",
    "policy": {"currency": "TWD", "tax_rate": "0.05"},
    "items": [
        {
            "id": "belt",
            "kind": "module",
            "subsystem": "輸送",
            "name": "輸送段",
            "unit": "組",
            "price": {"amount": "180000"},
        },
        {
            "id": "hmi",
            "kind": "purchased",
            "subsystem": "控制",
            "name": "HMI",
            "unit": "台",
            "price": {"amount": None},
        },
    ],
    "equipment": [
        {"id": "EQ-belt", "match": {"part": "library/conveyor.py"}, "items": [{"item": "belt"}]}
    ],
    "bom": [{"id": "B-hmi", "item": "hmi", "quantity": 1}],
}


def test_each_version_keeps_its_own_costing_for_api_and_export(tmp_path):
    project = create_project(
        tmp_path / "cost",
        {"name": "成本版本", "created": date.today().isoformat()},
        seed_example="getac_qc",
    )
    dump_yaml(project / "costing.yaml", GETAC_COSTING)
    v1 = build_project(project, "L0")
    assert v1["costing"]["complete"] is False and v1["costing"]["missing"] == 1
    changed = copy.deepcopy(GETAC_COSTING)
    changed["items"][0]["price"]["amount"] = "200000"
    dump_yaml(project / "costing.yaml", changed)
    v2 = build_project(project, "L0")
    conveyors = sum(
        1
        for machine in load_yaml(project / "cell.yaml")["machines"]
        for module in machine["modules"]
        if module.get("part") == "library/conveyor.py"
    )
    first = json.loads(
        (project / ".cellforge" / v1["version_id"] / "costing.json").read_text("utf-8")
    )
    second = json.loads(
        (project / ".cellforge" / v2["version_id"] / "costing.json").read_text("utf-8")
    )
    assert Decimal(second["totals"]["equipment"]["value"]) - Decimal(
        first["totals"]["equipment"]["value"]
    ) == conveyors * Decimal("20000")

    exported = export_project(project, None, v1["version_id"])
    files = {Path(name).suffix: Path(exported["directory"]) / name for name in exported["files"]}
    workbook = load_workbook(files[".xlsx"])
    assert {"摘要", "採購明細", "工時", "缺價清單", "版本"} <= set(workbook.sheetnames)
    detail = workbook["採購明細"]
    assert "ROUND(I2*J2*L2" in str(detail["M2"].value)
    assert detail["J2"].value == float(first["purchase_lines"][0]["unit_price"])
    assert workbook["缺價清單"]["B2"].value == "hmi"
    with next(
        path for name, path in files.items() if name == ".csv" and "purchase" in path.name
    ).open(encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    assert {row["採購項"] for row in rows} == {"belt", "hmi"}

    with TestClient(create_app(tmp_path, settings={"engineering_agent_mode": "local"})) as client:
        api_v1 = client.get(f"/api/projects/cost/versions/{v1['version_id']}/costing").json()
        api_v2 = client.get(f"/api/projects/cost/versions/{v2['version_id']}/costing").json()
        assert api_v1["totals"] == first["totals"] and api_v2["totals"] == second["totals"]

    plain = create_project(
        tmp_path / "plain",
        {"name": "無成本", "created": date.today().isoformat()},
        seed_example="getac_qc",
    )
    report = build_project(plain, "L0")
    assert report["costing"] is None
    assert "costing" not in export_project(plain, ["bom"], report["version_id"])["files"]
    with TestClient(create_app(tmp_path, settings={"engineering_agent_mode": "local"})) as client:
        assert (
            client.get(f"/api/projects/plain/versions/{report['version_id']}/costing").status_code
            == 404
        )
