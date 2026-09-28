"""Customer delivery-pack exporters."""

from __future__ import annotations

import base64
import csv
import html as html_lib
import json
import re
import shutil
import subprocess
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Any

import ezdxf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from cellforge import derived
from cellforge.project import source_root
from cellforge.versioning import MANIFEST_NAME, VersionContext, VersionError

ALL_KINDS = ("step", "dxf", "bom", "report", "deck", "html", "video")
# 成本表只在版本含 costing.json 時屬於「全部」；明確指定時缺成本資料會列為缺件。
COST_KINDS = ("costing", "purchase_bom")
NOT_EVALUATED = "未評估"
NOT_EVALUATED_DETAIL = (
    "此版本未執行 L1 工程檢查（沒有 checks.json），檢查結果為未評估，不代表通過。"
)


def export_project(
    project: Path, kinds: list[str] | None = None, version: str | None = None
) -> dict[str, Any]:
    """Export one version; kinds whose data that version lacks are listed, never substituted."""

    requested_all = not kinds
    kinds = [item.lower() for item in (kinds or list(ALL_KINDS))]
    unknown = sorted(set(kinds) - {*ALL_KINDS, *COST_KINDS})
    if unknown:
        raise ValueError(f"Unknown export kinds: {', '.join(unknown)}")
    context = VersionContext.open(project, version)
    if requested_all and context.optional_artifact("costing.json") is not None:
        kinds.extend(COST_KINDS)
    project_data = context.source_yaml("project.yaml", "取得專案名稱")
    target = project / "export" / context.name
    target.mkdir(parents=True, exist_ok=True)
    stem = f"{_safe_name(project_data['name'])}_{context.name}"
    producers: dict[str, tuple[str, Callable[[Path], dict[str, Any] | None]]] = {
        "step": (".step", lambda path: _copy(context.artifact("scene.step", "匯出 STEP"), path)),
        "dxf": ("_layout.dxf", lambda path: _dxf(context, path)),
        "bom": ("_bom.csv", lambda path: _bom(context, path)),
        "report": ("_report.docx", lambda path: _report(context, path)),
        "deck": ("_review_final3.pptx", lambda path: _deck(context, path)),
        "html": ("_review.html", lambda path: _html(context, path)),
        "video": ("_1080p.mp4", lambda path: _video(context, path)),
        "costing": ("_costing.xlsx", lambda path: _costing(context, path, "xlsx")),
        "purchase_bom": ("_purchase_bom.csv", lambda path: _costing(context, path, "csv")),
    }
    outputs: list[dict[str, Any]] = []
    missing: list[dict[str, str]] = []
    for kind in kinds:
        suffix, produce = producers[kind]
        output = target / f"{stem}{suffix}"
        try:
            provenance = produce(output)
        except VersionError as error:
            # 移除先前匯出的同名檔，避免下載到其他版本或已失效的內容。
            output.unlink(missing_ok=True)
            missing.append({"kind": kind, "reason": str(error)})
            continue
        outputs.append(
            {
                "kind": kind,
                "file": output.name,
                "sha256": derived.sha256_file(output),
                "size": output.stat().st_size,
                **({"provenance": provenance} if provenance else {}),
            }
        )
    manifest_path = context.directory / MANIFEST_NAME
    manifest = {
        "version": context.name,
        "status": "incomplete" if missing else "ok",
        "generated": datetime.now().astimezone().isoformat(),
        "version_complete": not context.legacy,
        "version_manifest_sha256": derived.sha256_file(manifest_path)
        if manifest_path.is_file()
        else None,
        "kinds": kinds,
        "files": [item["file"] for item in outputs],
        "outputs": outputs,
        "missing": missing,
        "warnings": context.warnings,
        "template": {
            "requested": "zq-work-deck",
            "used": "CellForge Midnight",
            "deviation": "zq-work-deck not supplied in workspace",
        },
    }
    (target / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return {
        "status": manifest["status"],
        "version": context.name,
        "directory": str(target),
        "files": manifest["files"],
        "missing": missing,
        "warnings": context.warnings,
    }


def _costing(context: VersionContext, target: Path, fmt: str) -> dict[str, Any]:
    from cellforge.costing.export import write_costing_xlsx, write_purchase_csv

    costing = context.artifact_json("costing.json", "匯出成本表")
    if fmt == "csv":
        write_purchase_csv(costing, target)
    else:
        manifest_path = context.directory / MANIFEST_NAME
        sources = {
            item.get("path"): item.get("sha256")
            for item in (context.manifest or {}).get("sources", [])
        }
        project = context.source_yaml("project.yaml", "匯出成本表")
        write_costing_xlsx(
            costing,
            target,
            {
                "version": context.name,
                "project_name": project.get("name", ""),
                "as_of": costing["as_of"],
                "complete": costing["complete"],
                "version_manifest_sha256": (
                    derived.sha256_file(manifest_path) if manifest_path.is_file() else ""
                ),
                "costing_yaml_sha256": sources.get("costing.yaml", ""),
                "exported": datetime.now().astimezone().isoformat(),
            },
        )
    return {"complete": costing["complete"], "missing": len(costing["missing"])}


def _video(context: VersionContext, target: Path) -> dict[str, Any]:
    artifact = derived.latest_available(context.project, context.name, "video")
    if artifact is None:
        raise derived.DerivedArtifactMissingError(
            derived.missing_reason(context.project, context.name, "video")
        )
    shutil.copy2(artifact.path, target)
    return {
        "derived_id": artifact.entry["id"],
        "path": artifact.entry["path"],
        "sha256": artifact.entry["sha256"],
        "settings": artifact.entry.get("settings", {}),
    }


def _copy(source: Path, target: Path) -> None:
    shutil.copy2(source, target)


def _dxf(context: VersionContext, target: Path) -> None:
    cell = context.source_yaml("cell.yaml", "匯出 DXF 平面配置")
    document = ezdxf.new("R2018", setup=True)
    document.units = ezdxf.units.MM
    model = document.modelspace()
    model.add_text("CellForge plan — dimensions in mm", dxfattribs={"height": 80}).set_placement(
        (-3800, 1800)
    )
    for machine in cell["machines"]:
        for module in machine["modules"]:
            x, y, _ = module.get("pose", {}).get("xyz", [0, 0, 0])
            params = module.get("params", {})
            length = float(params.get("length_mm", params.get("size", [500, 500])[0]))
            width = float(params.get("width_mm", params.get("size", [500, 500])[1]))
            points = [
                (x - length / 2, y - width / 2),
                (x + length / 2, y - width / 2),
                (x + length / 2, y + width / 2),
                (x - length / 2, y + width / 2),
            ]
            model.add_lwpolyline(points, close=True, dxfattribs={"layer": "EQUIPMENT"})
            model.add_text(module["id"], dxfattribs={"height": 55}).set_placement(
                (x - length / 2, y + width / 2 + 30)
            )
    document.saveas(target)


def _bom(context: VersionContext, target: Path) -> None:
    cell = context.source_yaml("cell.yaml", "匯出 BOM")
    modules = [module for machine in cell["machines"] for module in machine["modules"]]
    purpose = "匯出含廠商設備的 BOM"
    if any(module.get("vendor") for module in modules):
        vendor = context.source_yaml("vendor/manifest.yaml", purpose) or {}
    else:
        vendor = context.optional_source_yaml("vendor/manifest.yaml", purpose, {})
    vendor_map = {item["id"]: item for item in vendor.get("vendors", [])}
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["item", "module_id", "source", "vendor", "part_no", "trust", "quantity"],
        )
        writer.writeheader()
        for index, module in enumerate(modules, 1):
            item = vendor_map.get(module.get("vendor", ""), {})
            writer.writerow(
                {
                    "item": index,
                    "module_id": module["id"],
                    "source": module.get("part") or module.get("vendor"),
                    "vendor": module.get("vendor", ""),
                    "part_no": item.get("part_no", "TBD"),
                    "trust": module.get("trust") or module.get("pose", {}).get("trust", "inferred"),
                    "quantity": 1,
                }
            )


def _report(context: VersionContext, target: Path) -> None:
    purpose = "產生工程報告"
    data = context.source_yaml("project.yaml", purpose)
    assumptions = context.optional_source_yaml("analysis/assumptions.yaml", purpose, {}).get(
        "assumptions", []
    )
    checks = _optional_checks(context)
    timeline = context.artifact_json("timeline.json", purpose)
    step = context.artifact_json("step_validation.json", purpose)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    styles = document.styles
    styles["Normal"].font.name = "Microsoft JhengHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    styles["Normal"].font.size = Pt(9.5)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(data["name"])
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(31, 78, 121)
    subtitle = document.add_paragraph(f"Automated QC cell engineering report • {context.name}")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for warning in context.warnings:
        document.add_paragraph(warning)
    approximated = [
        str(module["id"])
        for module in (context.manifest or {}).get("modules", [])
        if module.get("approximated")
    ]
    if approximated:
        document.add_paragraph(
            f"近似廠商模型（非原廠真機）：{'、'.join(approximated)}；"
            "相關可達、干涉與關節速度檢查以型錄尺寸近似計算，交付前需以原廠模型確認。"
        )
    document.add_heading("Assumptions requiring confirmation", level=1)
    for item in assumptions:
        document.add_paragraph(
            f"{item['id']} — {item['text']} ({item.get('status', 'active')})", style="List Bullet"
        )
    document.add_page_break()
    document.add_heading("Executive engineering summary", level=1)
    document.add_paragraph(data.get("description", ""))
    table = document.add_table(rows=1, cols=4)
    table.style = "Light Shading Accent 1"
    for cell, value in zip(table.rows[0].cells, ("Metric", "Red", "Yellow", "Green"), strict=True):
        cell.text = value
    cells = table.add_row().cells
    summary = checks["summary"] if checks else {}
    for cell, value in zip(
        cells,
        (
            "L1 checks",
            summary.get("red", NOT_EVALUATED),
            summary.get("yellow", NOT_EVALUATED),
            summary.get("green", NOT_EVALUATED),
        ),
        strict=True,
    ):
        cell.text = str(value)
    document.add_paragraph(f"Timeline: {timeline['duration_s']:.1f} s at {timeline['fps']} fps")
    document.add_paragraph(
        f"STEP: {step['top_level_part_count']} top-level components / "
        f"{step['leaf_part_count']} leaves; names preserved = "
        f"{step['all_names_preserved']}"
    )
    images = _review_snapshots(context, checks)
    if images:
        image_table = document.add_table(rows=1, cols=len(images))
        for cell, image in zip(image_table.rows[0].cells, images, strict=True):
            cell.paragraphs[0].add_run().add_picture(str(image), width=Inches(3.25))
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_heading("Check results", level=1)
    if checks is None:
        document.add_paragraph(NOT_EVALUATED_DETAIL)
    for item in checks["items"] if checks else []:
        document.add_heading(f"{item['id']} • {item['severity'].upper()} • {item['type']}", level=2)
        document.add_paragraph(item.get("detail", ""))
        if item.get("suggestion"):
            document.add_paragraph(f"Action: {item['suggestion']}")
        document.add_paragraph(f"Evidence: {item.get('source', 'n/a')}")
    document.add_heading("Checklist mapping", level=1)
    checklist = context.optional_source("analysis/checklist_map.md", purpose)
    document.add_paragraph(
        checklist.read_text("utf-8") if checklist else "No checklist map supplied."
    )
    footer = section.footer.paragraphs[0]
    footer.text = (
        "CellForge • engineering evidence pack • inferred values require commissioning confirmation"
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(target)


def _deck(context: VersionContext, target: Path) -> None:
    purpose = "產生審查簡報"
    data = context.source_yaml("project.yaml", purpose)
    assumptions = context.optional_source_yaml("analysis/assumptions.yaml", purpose, {}).get(
        "assumptions", []
    )
    process = context.source_yaml("process.yaml", purpose)
    workpiece = context.source_yaml("workpiece.yaml", purpose)
    checks = _optional_checks(context)
    timeline = context.artifact_json("timeline.json", purpose)
    step = context.artifact_json("step_validation.json", purpose)
    snapshot = next(iter(_review_snapshots(context, checks)), None)
    payload = {
        "name": data["name"],
        "customer": data.get("customer", ""),
        "product": data.get("product", ""),
        "version": context.name,
        "assumptions": [
            f"{item['id']}: {item['text'][:45].rstrip()}{'...' if len(item['text']) > 45 else ''}"
            for item in assumptions
        ],
        "stations": [f"{item['id']} — {item['name']}" for item in process["stations"]],
        "checks": checks["summary"]
        if checks
        else {"red": NOT_EVALUATED, "yellow": NOT_EVALUATED, "green": NOT_EVALUATED},
        "checkItems": [
            f"[{item['severity'].upper()}] {item['id']}: "
            f"{item.get('value', 'n/a')} {item.get('unit', '')} / limit {item.get('limit', 'n/a')}"
            for item in checks["items"]
        ]
        if checks
        else [NOT_EVALUATED_DETAIL],
        "duration": timeline["duration_s"],
        "takt": process["takt"]["target_s"],
        "coverage": (
            f"{len([n for n in timeline['nodes'] if n.startswith('workpiece.')])}/"
            f"{len((workpiece.get('skus') or [{}])[0].get('covers', []))}"
        ),
        "step": step,
        "snapshot": str(snapshot.resolve()) if snapshot else "",
    }
    data_path = target.with_suffix(".deck-data.json")
    data_path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    script = source_root() / "tools" / "build_deck.ps1"
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    try:
        subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script),
                "-DataPath",
                str(data_path),
                "-OutputPath",
                str(target.resolve()),
            ],
            check=True,
            timeout=90,
            capture_output=True,
            creationflags=flags,
        )
    finally:
        data_path.unlink(missing_ok=True)


def _html(context: VersionContext, target: Path) -> dict[str, Any]:
    purpose = "產生離線 3D 審查頁"
    scene = base64.b64encode(context.artifact("scene.glb", purpose).read_bytes()).decode("ascii")
    timeline = context.artifact_json("timeline.json", purpose)
    checks = _optional_checks(context) or {"summary": {}, "items": []}
    title = str(context.source_yaml("project.yaml", purpose)["name"])
    bundle_path = _offline_bundle()
    bundle = bundle_path.read_text("utf-8").replace("</script", "<\\/script")
    payload = json.dumps(
        {"scene": scene, "timeline": timeline, "checks": checks, "title": title},
        ensure_ascii=False,
        separators=(",", ":"),
    ).replace("</script", "<\\/script")
    shell = (Path(__file__).resolve().parent / "_offline_viewer.html").read_text("utf-8")
    document = (
        shell.replace("__TITLE__", html_lib.escape(title))
        .replace("__VERSION__", context.name)
        .replace("__PAYLOAD__", payload)
        .replace("__BUNDLE__", bundle)
    )
    target.write_text(document, encoding="utf-8")
    return {"offline_viewer_bundle_sha256": derived.sha256_file(bundle_path)}


def _offline_bundle() -> Path:
    candidates = (
        source_root() / "web" / "dist" / "offline.js",
        Path(__file__).resolve().parent / "_viewer" / "offline.js",
        Path(__file__).resolve().parent / "_offline_viewer.js",
    )
    bundle_path = next((path for path in candidates if path.is_file()), None)
    if bundle_path is None:
        raise FileNotFoundError("缺少離線 3D viewer bundle；請先在 web 執行 npm run build")
    return bundle_path


def _optional_checks(context: VersionContext) -> dict[str, Any] | None:
    """L0 版本沒有 checks.json：回傳 None，由呼叫端標示「未評估」，不可當成通過。"""
    path = context.optional_artifact("checks.json")
    return json.loads(path.read_text("utf-8")) if path else None


def _review_snapshots(context: VersionContext, checks: dict[str, Any] | None) -> list[Path]:
    # 舊版快照曾把截圖直接放進 vN；新版截圖是登記在該版本名下的衍生檔。
    images = list(context.directory.glob("snapshot_t*.png")) + [
        item.path
        for item in derived.artifacts(context.project, context.name, "snapshot")
        if item.available
    ]
    candidates: dict[Path, float] = {}
    for path in images:
        match = re.match(r"snapshot_t(-?\d+(?:\.\d+)?)_", path.name)
        if match and path not in candidates:
            candidates[path] = float(match.group(1))
    desired = []
    for severity in ("red", "yellow"):
        timed = [
            item
            for item in (checks or {}).get("items", [])
            if item.get("severity") == severity and isinstance(item.get("t"), int | float)
        ]
        if timed:
            desired.append(float(max(timed, key=_check_badness)["t"]))
    selected: list[Path] = []
    for check_time in desired:
        available = [path for path in candidates if path not in selected]
        if available:
            selected.append(min(available, key=lambda path: abs(candidates[path] - check_time)))
    if selected:
        return selected
    by_name = {path.name: path for path in reversed(images)}
    overview = ("snapshot_t0_iso.png", "snapshot_station_s3.png")
    return [by_name[name] for name in overview if name in by_name]


def _check_badness(item: dict[str, Any]) -> float:
    distance = item.get("min_dist_mm")
    if isinstance(distance, int | float):
        return -float(distance)
    value, limit = item.get("value"), item.get("limit")
    if isinstance(value, int | float) and isinstance(limit, int | float):
        return abs(float(value) - float(limit)) / max(abs(float(limit)), 1e-9)
    return 0.0


def _safe_name(value: str) -> str:
    return (
        "".join(char if char.isalnum() or char in "-_" else "_" for char in value).strip("_")
        or "cellforge"
    )
