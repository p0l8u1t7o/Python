"""
產生 docs/cost-estimate.xlsx（鋁質殼體 AOI＋共焦量測半自動設備，預算級估價）並印出摘要數字給 cost-estimate.md。

    python tools/build_cost_estimate.py

藍字為可改的單價、數量、費率；數量 0 的選配不計入。金額為內部預算估計，非供應商報價。
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parents[1] / "docs" / "cost-estimate.xlsx"

ENG_RATE, TECH_RATE, CONTINGENCY, TAX = 8000, 5500, 0.15, 0.05
GRADES = {"A": (0.10, "型錄價或近期同級採購價"), "B": (0.20, "同級品預算價，需詢價"), "C": (0.30, "自製或規格未定，幅度較大")}

# (編號, 子系統, 項目, 規格要求, 建議選型, 選型理由, 數量, 單位, 單價, 等級)
ITEMS = [
    ("1-01", "1 ST1 光學", "通道 A 線掃相機", "4K 彩色線掃、7 µm 圖元、線頻 ≥ 13 kHz", "Basler raL4096-24gc 或同級", "一次掃描取得 RGB 三角度影像", 1, "台", 160000, "B"),
    ("1-02", "1 ST1 光學", "通道 A 線掃遠心鏡頭", "1.5×、像面 ≥ 28.7 mm、含同軸介面", "線掃遠心 1.5×", "視野 19.1 mm 涵蓋規格 C 全長", 1, "支", 140000, "B"),
    ("1-03", "1 ST1 光學", "RGB 三角度線光源", "紅 15° 暗場、綠 60° 明場、藍同軸＋三通道恆流控制器", "CCS／OPT 線光源", "拉傷、凹陷、污染各用最適照明", 1, "組", 120000, "B"),
    ("1-04", "1 ST1 光學", "通道 B 相機＋雙遠心鏡頭", "5 MP 全域快門 GigE、0.5× 雙遠心", "IDS GigE 5MP＋0.5× 雙遠心", "背光剪影量外徑／全長，次圖元 ±1 µm", 1, "組", 135000, "B"),
    ("1-05", "1 ST1 光學", "通道 B 遠心平行背光", "525 nm、口徑 ≥ 30 mm、頻閃", "遠心平行背光", "與遠心鏡頭成對才有銳利邊緣", 1, "組", 45000, "B"),
    ("1-06", "1 ST1 光學", "通道 C 相機＋遠心鏡頭", "5 MP GigE、1× 遠心含同軸光、環形低角度光", "IDS GigE 5MP＋1× 遠心", "由中空軸向下看杯口端面", 1, "組", 120000, "B"),
    ("1-07", "1 ST1 光學", "光學支架與調整座", "懸臂立柱、三軸微調、遮光罩", "自製", "光學件全部懸臂，讓出後側移載走廊", 1, "式", 130000, "C"),
    ("2-01", "2 ST1 機構", "DD 中空軸馬達＋驅動器", "中空孔 ≥ Ø20、徑跳 ≤ 2 µm、24 bit 編碼器", "DD 馬達＋EtherCAT 驅動器", "編碼器硬體觸發線掃", 1, "組", 180000, "B"),
    ("2-02", "2 ST1 機構", "三爪 PEEK 夾頭", "夾持帶 1.5 mm、限力 2 N、夾持到位感測；三規格爪各一組", "自製＋微型氣動夾頭＋精密減壓閥", "0.18 mm 薄壁不可變形", 1, "組", 150000, "C"),
    ("2-03", "2 ST1 機構", "立柱與頭部座", "鑄鐵或鋼構立柱、馬達座", "自製", "", 1, "式", 90000, "C"),
    ("3-01", "3 ST2 量測", "下共焦感測器", "量程 1 mm、線性度 ±0.25 µm、encoder 觸發、EtherCAT", "Micro-Epsilon confocalDT IFD2410-1", "外底面無遮擋，選精度最高", 1, "台", 420000, "B"),
    ("3-02", "3 ST2 量測", "上共焦感測器（規格 A／B）", "量程 6 mm、NA 0.18、線性度 ±1.5 µm", "Micro-Epsilon confocalDT IFD2410-6", "光束可通過 Ø4.26／Ø4.42 杯口", 1, "台", 420000, "B"),
    ("3-03", "3 ST2 量測", "上共焦（規格 C）控制器＋光纖探頭", "小 NA 長量程探頭、量程 ≥ 20 mm", "confocalDT IFC2411＋IFS 探頭", "IFD2410 全系列被 Ø3.72 內孔遮擋", 1, "組", 480000, "B"),
    ("3-04", "3 ST2 量測", "空心軸 θ 平台", "通孔 Ø58、徑跳 < 1 µm、編碼器輸出", "DD 空心旋轉平台＋驅動器", "下感測器伸入通孔", 1, "組", 260000, "B"),
    ("3-05", "3 ST2 量測", "R 軸微動台＋Z 調整台", "R 行程 ±3 mm 解析度 0.1 µm、Z 行程 20 mm", "精密電動滑台", "R 軸承載整個 C 型架", 1, "組", 170000, "B"),
    ("3-06", "3 ST2 量測", "C 型架與橋板", "鋼構應力消除、上下感測器共線調整", "自製", "", 1, "式", 110000, "C"),
    ("3-07", "3 ST2 量測", "可換式環座治具", "鎢鋼薄環座平面度 ≤ 1 µm、定位銷 Ø6 h6；A／B／C 各 1＋備品 1", "自製（鎢鋼研磨）", "換型 60 秒內完成", 4, "套", 45000, "C"),
    ("4-01", "4 移載與托盤", "X 軸", "行程 520 mm、重複精度 ±5 µm、1.2 m/s", "線性馬達或滾珠螺桿模組", "托盤到 ST2 一支軸", 1, "組", 140000, "B"),
    ("4-02", "4 移載與托盤", "Z 軸電動滑台", "行程 80 mm", "小型電動滑台", "三規格高度不同，不用氣缸", 1, "組", 60000, "B"),
    ("4-03", "4 移載與托盤", "貼靠氣缸＋側向真空吸嘴", "行程 6 mm、V 形導電吸嘴、真空開關", "SMC MXQ／ZK2 或同級＋自製吸嘴", "吸外壁，不碰杯口與底面", 1, "組", 55000, "C"),
    ("4-04", "4 移載與托盤", "托盤梭台", "入料、出料各一軸，行程 270 mm", "電動滑台", "對列＋送到前方上下料位", 2, "組", 70000, "B"),
    ("4-05", "4 移載與托盤", "導電托盤", "4×5 與 4×2，A／B／C 各一套，穴深 1 mm", "自製（導電 POM）", "防靜電吸塵", 12, "片", 6000, "C"),
    ("5-01", "5 機台與控制", "花崗岩平台＋避震", "620 × 430 × 42 mm、四點避震腳座", "花崗岩平台", "熱膨脹低、阻尼好", 1, "式", 120000, "B"),
    ("5-02", "5 機台與控制", "底櫃與外罩", "鋼構底櫃、鋁擠外罩、量測區前門互鎖", "自製", "", 1, "式", 260000, "C"),
    ("5-03", "5 機台與控制", "PLC 與 EtherCAT 主站", "KV-X＋KV-XH16EC＋I/O", "Keyence", "與其他專案相同平台", 1, "套", 220000, "A"),
    ("5-04", "5 機台與控制", "安全控制器與門鎖", "GC-1000、安全門鎖、急停", "Keyence", "門開切斷運動", 1, "套", 90000, "A"),
    ("5-05", "5 機台與控制", "視覺工控機", "RTX 4060 Ti 級 GPU、影像擷取卡、GigE PoE 網卡", "工業電腦", "TensorRT FP16 複判", 1, "套", 190000, "B"),
    ("5-06", "5 機台與控制", "HMI 與燈號", "15 吋觸控螢幕、三色燈", "工業觸控螢幕", "", 1, "式", 55000, "A"),
    ("5-07", "5 機台與控制", "電控與氣動", "電控盤、配線、調壓、真空產生器", "自製", "", 1, "式", 180000, "C"),
    ("6-01", "6 校正件", "標準厚度片", "0.15／0.30／1.00 mm，U ≤ 0.3 µm，附校正報告", "校正實驗室出證", "每班校正差分公式 K 值", 3, "片", 25000, "B"),
    ("6-02", "6 校正件", "標準環規與標定板", "Ø4.80 h4 環規、圓點陣列標定板、光學平面片", "外購", "尺寸驗證與畸變標定", 1, "組", 90000, "B"),
    ("7-01", "7 工程", "機構設計", "光學支架、夾頭、治具、移載、機台", "—", "", 45, "人日", ENG_RATE, "B"),
    ("7-02", "7 工程", "電控設計", "電路、PLC、觸發同步、安全回路", "—", "", 25, "人日", ENG_RATE, "B"),
    ("7-03", "7 工程", "視覺與量測軟體", "採集、展開、尺寸量測、點雲、配方、資料庫、HMI", "—", "", 70, "人日", ENG_RATE, "B"),
    ("7-04", "7 工程", "深度學習複判", "黃金樣本庫、標註、訓練、部署", "—", "", 30, "人日", ENG_RATE, "C"),
    ("7-05", "7 工程", "POC", "光錐可及性、照明對比、夾持變形實測", "—", "", 20, "人日", ENG_RATE, "B"),
    ("7-06", "7 工程", "組立配線", "機構組立、配線、配管", "—", "", 30, "人日", TECH_RATE, "B"),
    ("7-07", "7 工程", "校正、GR&R 與驗收", "幾何校正、MSA、連續運轉", "—", "", 25, "人日", ENG_RATE, "B"),
    ("8-01", "8 選配", "內壁檢測", "孔內檢測鏡頭＋內置環形照明＋相機", "Opto Engineering PCHI 或同級", "R-06：內壁若列為必檢", 0, "組", 260000, "B"),
    ("8-02", "8 選配", "自動上下料", "SCARA＋托盤抽屜，取代人工擺盤", "DENSO SCARA", "產能需求超過 500 UPH 時", 0, "式", 1400000, "C"),
    ("8-03", "8 選配", "接觸式深度計（規格 C 備案）", "LVDT＋Ø1.5 探針、測力 < 0.5 N", "induSENSOR DTD 或同級", "光纖探頭不可行時", 0, "組", 120000, "B"),
]

HEAD = ["編號", "子系統", "項目", "規格要求", "建議選型", "選型理由", "數量", "單位", "單價 NT$", "小計 NT$", "等級", "下限 NT$", "上限 NT$", "類別"]
BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="C9CED6"))
FILL = PatternFill("solid", fgColor="E8EEF6")
MONEY = '#,##0'


def main() -> None:
    wb = Workbook()
    # ---- 明細 ----
    ws = wb.active; ws.title = "明細"
    ws.append(HEAD)
    for c in ws[1]: c.font = BOLD; c.fill = FILL
    for i, it in enumerate(ITEMS, start=2):
        code, sub, name, spec, model, why, qty, unit, price, grade = it
        ws.append([code, sub, name, spec, model, why, qty, unit, price, f"=G{i}*I{i}", grade,
                   f"=J{i}*(1-VLOOKUP(K{i},摘要!$A$12:$B$14,2,FALSE))", f"=J{i}*(1+VLOOKUP(K{i},摘要!$A$12:$B$14,2,FALSE))", int(code.split("-")[0])])
        for col in "GI": ws[f"{col}{i}"].font = BLUE
        for col in "IJLM": ws[f"{col}{i}"].number_format = MONEY
    n = len(ITEMS) + 1
    widths = [7, 14, 20, 40, 26, 30, 7, 6, 12, 13, 6, 13, 13, 6]
    for k, w in enumerate(widths, start=1): ws.column_dimensions[get_column_letter(k)].width = w
    ws.freeze_panes = "A2"

    # ---- 摘要 ----
    s = wb.create_sheet("摘要", 0)
    rows = [
        ["鋁質殼體 AOI＋共焦量測半自動設備 預算級估價"],
        ["估算日 2026-09-29。單套全新設備、新台幣、未取得供應商報價。"],
        ["單機一套；人工擺盤、機內自動取放；ST1 三通道光學＋ST2 上下共焦差分測厚；涵蓋規格 A／B／C。"],
        ["參數（藍字可改）"],
        ["工程費率（NT$/人日）", ENG_RATE, "設計、軟體、POC 與驗收"],
        ["技術費率（NT$/人日）", TECH_RATE, "組立配線"],
        ["預備費率", CONTINGENCY, "套用設備＋工程＋已選選配"],
        ["營業稅率（預算假設）", TAX, "含稅＝含預備費總額 × (1＋稅率)"],
        [],
        ["估價等級與幅度"],
        ["等級", "幅度", "定義"],
    ]
    for r in rows: s.append(r)
    for g, (rng, txt) in GRADES.items(): s.append([g, rng, txt])
    s.append([])
    s.append(["項目", "金額 NT$", "說明"])
    base = s.max_row
    s.append(["設備與材料（1～6 類）", f'=SUMIFS(明細!$J$2:$J${n},明細!$N$2:$N${n},"<7")', "不含工程人日與選配"])
    s.append(["工程人日（7 類）", f'=SUMIFS(明細!$J$2:$J${n},明細!$N$2:$N${n},7)', "設計、程式、組立與驗收"])
    s.append(["已選選配（8 類）", f'=SUMIFS(明細!$J$2:$J${n},明細!$N$2:$N${n},8)', "數量改為 1 才計入"])
    s.append(["小計（未稅）", f"=SUM(B{base+1}:B{base+3})", ""])
    s.append(["預備費", f"=B{base+4}*B7", ""])
    s.append(["含預備費，未稅", f"=B{base+4}+B{base+5}", "預算主數字"])
    s.append(["下限（含預備費，未稅）", f"=SUM(明細!L2:L{n})*(1+B7)", "逐列幅度加總，不是統計區間"])
    s.append(["上限（含預備費，未稅）", f"=SUM(明細!M2:M{n})*(1+B7)", ""])
    s.append(["含稅預算", f"=B{base+6}*(1+B8)", ""])
    for r in range(base + 1, base + 10): s[f"B{r}"].number_format = MONEY
    for cell in ("B5", "B6", "B7", "B8"): s[cell].font = BLUE
    s["B7"].number_format = s["B8"].number_format = "0%"
    for r in range(12, 15): s[f"B{r}"].number_format = "0%"
    s["A1"].font = Font(bold=True, size=14)
    for r in (4, 10, base): s[f"A{r}"].font = BOLD
    s.column_dimensions["A"].width = 26; s.column_dimensions["B"].width = 16; s.column_dimensions["C"].width = 60

    # ---- 選型計算 ----
    c = wb.create_sheet("選型計算")
    calc = [
        ["選型與數量核對"], [],
        ["項目", "數值", "單位", "依據"],
        ["通道 A 軸向解析度", "=7/1.5", "µm/px", "圖元 7 µm ÷ 倍率 1.5"],
        ["通道 A 每圈行數", "=PI()*4.9/(B4/1000)", "行", "Ø4.9 周長 ÷ 解析度"],
        ["通道 A 線頻（1 rev/s）", "=B5/1000", "kHz", "相機上限 24 kHz"],
        ["通道 B 影像解析度", "=3.45/0.5", "µm/px", "圖元 3.45 µm ÷ 倍率 0.5"],
        ["規格 A 杯口光束直徑", "=2*(4.2-0.26)*TAN(ASIN(0.18))", "mm", "內徑 Ø4.26，IFD2410-6"],
        ["規格 B 杯口光束直徑", "=2*(9.86-0.18)*TAN(ASIN(0.18))", "mm", "內徑 Ø4.42，單邊餘裕約 0.45"],
        ["規格 C 杯口光束直徑（NA 0.10）", "=2*(13.75-0.88)*TAN(ASIN(0.10))", "mm", "內孔 Ø3.72；NA 為推估，R-08 待供應商確認"],
        ["螺旋掃描點數", "=5*3600", "點", "5 圈 × 3600 點"],
        ["螺旋掃描時間", "=5*60/120", "s", "120 rpm"],
        ["規劃節拍（3D 模擬）", 10.4, "s／件", "單件流；規格書估 9.4 s，模擬多了 ST2 前的升降點"],
        ["每小時產能（稼動 85%）", "=3600/B13*0.85", "件/h", ""],
        ["底厚擴充不確定度 U（k=2）", 1.87, "µm", "規格書 3.7 節；U/T = 4.7 %"],
    ]
    for r in calc: c.append(r)
    c["A1"].font = Font(bold=True, size=13)
    for cell in c[3]: cell.font = BOLD
    for w, col in zip([24, 14, 8, 50], "ABCD"): c.column_dimensions[col].width = w
    for r in range(4, 17): c[f"B{r}"].number_format = "0.00"

    # ---- 通訊架構 ----
    t = wb.create_sheet("通訊架構")
    for r in [["通訊與整合範圍"], [], ["裝置", "連接", "介面／協定", "交換內容", "待確認"],
              ["θ1 編碼器", "線掃相機", "硬體 A/B 分頻", "行觸發 3300 行／圈", "分頻比與抖動"],
              ["θ2 編碼器", "上下共焦", "encoder 輸入（位置觸發）", "同角度取樣", "兩支感測器同源觸發"],
              ["EtherCAT 主站", "伺服、共焦", "EtherCAT＋DC", "運動、量測值、同步時基", "抖動 ≤ 1 µs"],
              ["PLC", "光源控制器、面陣相機", "硬體觸發", "頻閃與曝光同步", ""],
              ["視覺工控機", "相機", "GigE Vision／CoaXPress", "影像", "線掃介面"],
              ["視覺工控機", "PLC", "EtherNet/IP 或 TCP", "配方、結果、分流碼", ""],
              ["安全控制器", "門鎖、急停、伺服 STO", "安全 I/O", "門開切斷運動", "獨立安全回路"]]:
        t.append(r)
    t["A1"].font = Font(bold=True, size=13)
    for cell in t[3]: cell.font = BOLD
    for w, col in zip([12, 22, 22, 30, 24], "ABCDE"): t.column_dimensions[col].width = w

    # ---- 假設與待確認 ----
    a = wb.create_sheet("假設與待確認")
    for r in [["估價假設與待確認"], [], ["項目", "內容"],
              ["估價基礎", "2026-09-29 內部預算估計，不是供應商報價；型號為建議，採購前需詢價與確認交期。"],
              ["範圍", "單機一套，涵蓋規格 A／B／C；含三規格環座、夾爪、托盤與校正件。"],
              ["不含", "樣品、廠務（電、氣、網路、控溫）改造、MES 介面開發、年度校正與保養合約、銷售毛利。"],
              ["待檢件", "以未裝填的空殼為前提（R-03）；若不是，本設計與估價都不適用。"],
              ["規格 C", "上感測器另購 IFC2411＋光纖探頭；NA 為推估值，需供應商實測規格（R-08）。"],
              ["外觀基準", "允收基準未定（R-05）；深度學習人日以 300 件樣品建庫估算。"],
              ["節拍", "3D 模擬規劃值約 10.4 s／件，不是驗收承諾。"],
              ["上下限", "逐列依等級幅度加總（A ±10%、B ±20%、C ±30%），不是統計信賴區間。"]]:
        a.append(r)
    a["A1"].font = Font(bold=True, size=13)
    for cell in a[3]: cell.font = BOLD
    a.column_dimensions["A"].width = 12; a.column_dimensions["B"].width = 90
    for row in a.iter_rows(min_row=4): row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.row > 1: cell.border = THIN
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    # ---- 同樣的算式，印出摘要數字 ----
    sub = lambda lo, hi: sum(q * p for code, *_, q, _u, p, _g in ITEMS if lo <= float(code.split('-')[0]) < hi)
    eq, eng, opt = sub(1, 7), sub(7, 8), sub(8, 9)
    subtotal = eq + eng + opt
    total = subtotal * (1 + CONTINGENCY)
    lo = sum(q * p * (1 - GRADES[g][0]) for *_, q, _u, p, g in ITEMS) * (1 + CONTINGENCY)
    hi = sum(q * p * (1 + GRADES[g][0]) for *_, q, _u, p, g in ITEMS) * (1 + CONTINGENCY)
    print(f"設備與材料 {eq:,.0f}\n工程人日 {eng:,.0f}\n選配 {opt:,.0f}\n小計 {subtotal:,.0f}\n含預備費未稅 {total:,.0f}\n下限 {lo:,.0f}\n上限 {hi:,.0f}\n含稅 {total * (1 + TAX):,.0f}\n項目數 {len(ITEMS)}")


if __name__ == "__main__":
    main()
